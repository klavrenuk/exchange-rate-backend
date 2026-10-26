import os
from datetime import date, datetime
from typing import Optional
from sqlalchemy import select, func
from sqlalchemy.ext.asyncio import AsyncSession
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from app.common.models import Rate, Currency
from app.core.config import settings
from app.services.email.sender import send_email_with_attachment


async def resolve_report_date(session: AsyncSession, report_date: Optional[date] = None) -> date:
    requested = report_date or date.today()
    res = await session.execute(select(Rate.id).where(Rate.date == requested).limit(1))
    if res.scalar_one_or_none() is not None:
        return requested
    res = await session.execute(select(func.max(Rate.date)).where(Rate.date <= requested))
    latest = res.scalar_one_or_none()
    return latest or requested


async def generate_excel_report(session: AsyncSession, report_date: Optional[date] = None):
    report_date = await resolve_report_date(session, report_date)
    os.makedirs(settings.REPORT_DIR, exist_ok=True)
    q = select(Rate, Currency).join(Currency, Rate.currency_id == Currency.id).where(Rate.date == report_date).order_by(Currency.abbreviation)
    res = await session.execute(q)
    rows = res.all()
    wb = Workbook()
    ws = wb.active
    ws.title = "Rates"
    ws["A1"] = f"Exchange rates for {report_date.isoformat()} (NBRB)"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:E1")
    ws["A2"] = "Date"
    ws["B2"] = "Currency"
    ws["C2"] = "Name"
    ws["D2"] = "Scale"
    ws["E2"] = "Rate"
    for cell in ["A2", "B2", "C2", "D2", "E2"]:
        ws[cell].font = Font(bold=True)
        ws[cell].alignment = Alignment(horizontal="center")
    r = 3
    for rate, cur in rows:
        ws.cell(row=r, column=1).value = rate.date
        ws.cell(row=r, column=2).value = cur.abbreviation
        ws.cell(row=r, column=3).value = cur.name
        ws.cell(row=r, column=4).value = rate.scale
        ws.cell(row=r, column=5).value = rate.rate
        r += 1
    for col in ["A", "B", "C", "D", "E"]:
        ws.column_dimensions[col].width = 18
    filename = f"nbrb_rates_{report_date.isoformat()}.xlsx"
    path = os.path.join(settings.REPORT_DIR, filename)
    wb.save(path)
    return path


async def report_and_send(session: AsyncSession, report_date: Optional[date] = None):
    from app.common.models import JobRun
    import json
    job = JobRun(job_name="nbrb.report_and_send", status="running")
    session.add(job)
    await session.commit()
    await session.refresh(job)
    try:
        actual_date = await resolve_report_date(session, report_date)
        path = await generate_excel_report(session, actual_date)
        subject = f"{settings.MAIL_SUBJECT_PREFIX} NBRB rates {actual_date.isoformat()}"
        body = f"Exchange rates report attached.\nDate: {actual_date.isoformat()}\nSource: National Bank of the Republic of Belarus (NBRB)"
        await send_email_with_attachment(subject, body, path)
        job.status = "success"
        job.finished_at = datetime.utcnow()
        if job.started_at:
            job.duration_ms = int((job.finished_at - job.started_at).total_seconds() * 1000)
        details = {"report_path": path, "report_date": actual_date.isoformat()}
        if report_date is not None and report_date != actual_date:
            details["requested_date"] = report_date.isoformat()
        job.details_json = json.dumps(details)
        await session.commit()
        return path
    except Exception as e:
        job.status = "failed"
        job.finished_at = datetime.utcnow()
        if job.started_at:
            job.duration_ms = int((job.finished_at - job.started_at).total_seconds() * 1000)
        job.error_text = str(e)
        await session.commit()
        raise
